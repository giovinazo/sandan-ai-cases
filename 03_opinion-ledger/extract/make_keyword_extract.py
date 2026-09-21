# -*- coding: utf-8 -*-
"""파생 추출: 특정 주제(예: 무상귀속·국공유지) 관련 협의의견만 뽑아 엑셀로 낸다.

판정(설정 extract)
  · A등급: 협의의견·조치계획 본문(공백 제거)에 키워드가 있는 건
  · B등급: 미회신 부서 중 소관부서명이 재산관리 계열인 건(회신 확인 필요)
  · pinned: 사람이 확정한 건을 연번 + 지문(소관부서 | 협의의견 앞 20자)으로 잠근다.
    판이 올라가면 연번이 밀리므로, 지문이 어긋나면 조용히 엉뚱한 건을 뽑지 않고 멈춘다.

실행: python3 extract/make_keyword_extract.py
"""
import json
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "pipeline"))
import cfg  # noqa: E402

X = cfg.C["extract"]
TEAL, FONT = "1F4E46", "맑은 고딕"
thin = Side(style="thin", color="D9D4C7")
BD = Border(left=thin, right=thin, top=thin, bottom=thin)


def flat(s):
    return re.sub(r"\s+", "", s or "")


def fingerprint(r):
    return f"{r['소관부서']}|{flat(r['협의의견'])[:20]}"


def pick(recs):
    by = {r["연번"]: r for r in recs}
    out = {}
    for p in X.get("pinned", []):                       # 사람이 확정한 건(지문 잠금)
        r = by.get(p["연번"])
        if not r or fingerprint(r) != p["지문"]:
            raise SystemExit(f"pinned 연번 {p['연번']} 지문 불일치. 판이 바뀌었으면 연번을 재매핑할 것")
        out[r["연번"]] = (p.get("등급", "A"), p.get("근거", "수기 확정"))
    kws = [flat(k) for k in X["keywords"]]
    for r in recs:
        if r["연번"] in out:
            continue
        body = flat(r["협의의견"] + r["조치계획"])
        hit = [k for k in kws if k in body]
        if hit:
            out[r["연번"]] = ("A", "키워드: " + "·".join(hit))
        elif r["반영여부"] == "미회신" and any(k in r["소관부서"] for k in X.get("dept_keywords_for_noreply", [])):
            out[r["연번"]] = ("B", "재산관리 계열 부서 미회신(회신 확인 필요)")
    return [(by[n], g, why) for n, (g, why) in sorted(out.items(), key=lambda kv: (kv[1][0], kv[0]))]


def main():
    recs = json.load(open(cfg.work("records.json"), encoding="utf-8"))
    rows = pick(recs)
    wb = Workbook()
    ws = wb.active
    ws.title = "추출"
    cols = ["등급", "연번", "기관군", "관계기관", "소관부서", "반영여부", "판정근거", "협의의견", "조치계획"]
    widths = [6, 6, 9, 14, 16, 9, 30, 60, 60]
    ws.append([f"■ {X['title']}  (기준일 {cfg.C['project']['base_date']} · {cfg.C['project']['edition']})"])
    ws["A1"].font = Font(name=FONT, size=12, bold=True, color=TEAL)
    ws.append(cols)
    for j, w in enumerate(widths, 1):
        c = ws.cell(2, j)
        c.font = Font(name=FONT, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=TEAL)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[c.column_letter].width = w
    for r, g, why in rows:
        ws.append([g, r["연번"], r["기관군"], r["관계기관"], r["소관부서"], r["반영여부"], why,
                   r["협의의견"], r["조치계획"]])
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.border = BD
            c.font = Font(name=FONT, size=9)
            c.alignment = Alignment(vertical="top", wrap_text=c.column_letter in ("G", "H", "I"))
    ws.freeze_panes = "A3"
    wb.save(cfg.out("extract"))
    a = sum(1 for _, g, _ in rows if g == "A")
    print(f"추출 {len(rows)}건 (A {a} · B {len(rows) - a}) → {cfg.out('extract').relative_to(cfg.ROOT)}")


if __name__ == "__main__":
    main()
