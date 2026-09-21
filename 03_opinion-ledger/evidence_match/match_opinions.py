# -*- coding: utf-8 -*-
"""대장 ↔ 회신 원문 공문 건별 대조(선택 기능).

대장 «협의의견»이 관계기관이 실제로 보낸 회신 공문에 있는지 건별로 대조한다.
요약본(용역사 정리표)만 믿지 않고 원문과 맞춰 보기 위한 장치다.

대조 원리: 양쪽 다 공백을 전부 걷어낸 뒤(OCR의 음절 띄어쓰기 대응) 문자 3그램 포함률로 잰다.
  포함률 = |3그램(대장 의견) ∩ 3그램(원문)| ÷ |3그램(대장 의견)|
  일치 ≥ 0.85 · 부분일치 ≥ 0.60 · 그 밑은 상이·확인불가(사람이 원문 확인)
요약·개행·오탈자에 관대하다. 어느 대목과 겹치는지는 최장 공통 부분열로 발췌한다.

입력: config/evidence_map.json(공문 파일 ↔ 기관·부서), paths.evidence_dir 안의 PDF
      글자가 없는 스캔 PDF는 macOS에서 vision_ocr.py(Apple Vision)로 읽는다(다른 OS는 건너뜀)
산출: output/work/match_results.json(HTML 빌더가 근거 공문 연결에 사용)
      output/<match_xlsx>(건별 대조표)

실행: python3 evidence_match/match_opinions.py
"""
import difflib
import json
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from pathlib import Path

import openpyxl
import pdfplumber
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "pipeline"))
import cfg  # noqa: E402

TRIV = re.compile(r"^[○\s\-·]*((의견|검토의견)\s?없음|해당\s?(사항)?\s?없음|저촉\s?(사항)?\s?없음"
                  r"|미회신|이상\s?없음|별도\s?의견\s?없음)")
NFC = lambda s: unicodedata.normalize("NFC", str(s if s is not None else "")).strip()


def norm_map(s):
    """정규화 문자열과 (정규화 위치 → 원본 위치) 배열"""
    s = unicodedata.normalize("NFC", s or "")
    out, idx = [], []
    for i, ch in enumerate(s):
        if ch.isspace():
            continue
        out.append(ch.lower())
        idx.append(i)
    return "".join(out), idx


def norm(s):
    return norm_map(s)[0]


def grams(s, n=3):
    return {s[i:i + n] for i in range(len(s) - n + 1)} if len(s) >= n else ({s} if s else set())


def coverage(op_n, doc_grams):
    g = grams(op_n)
    return len(g & doc_grams) / len(g) if g else 0.0


def read_text(path):
    """PDF 글자 추출. 글자가 거의 없으면(스캔본) macOS Vision OCR로 넘긴다."""
    with pdfplumber.open(str(path)) as p:
        text = "\n".join((pg.extract_text() or "") for pg in p.pages)
    if len(norm(text)) >= 30:
        return text, "pdf텍스트"
    try:
        import vision_ocr                                   # macOS 전용(pyobjc)
        return vision_ocr.ocr_pdf_file(str(path)), "ocr"
    except Exception as e:                                  # noqa: BLE001
        print(f"  ⚠ 스캔본 OCR 불가({type(e).__name__}): {path.name}")
        return text, "추출실패"


def load_docs():
    emap = json.load(open(cfg.path("evidence_map"), encoding="utf-8"))["문서"]
    base = cfg.path("evidence_dir")
    by_led = defaultdict(list)
    for m in emap:
        f = base / m["file"]
        if not f.exists():
            print(f"  ⚠ 원문 없음: {m['file']}")
            continue
        text, how = read_text(f)
        d = dict(m, text=text, 추출방식=how)
        d["_norm"], d["_idx"] = norm_map(text)
        d["_grams"] = grams(d["_norm"])
        by_led[(NFC(m["기관군"]), NFC(m["관계기관"]), NFC(m["소관부서"]))].append(d)
    return by_led


def excerpt(op_n, doc):
    """원문에서 의견과 가장 길게 겹치는 대목을 원본 표기로 발췌"""
    m = difflib.SequenceMatcher(None, op_n, doc["_norm"], autojunk=False)
    blk = m.find_longest_match(0, len(op_n), 0, len(doc["_norm"]))
    if blk.size < 8:
        return ""
    a, b = doc["_idx"][blk.b], doc["_idx"][blk.b + blk.size - 1]
    raw = unicodedata.normalize("NFC", doc["text"])[max(0, a - 15):b + 16]
    return re.sub(r"\s+", " ", raw).strip()


def verdict_of(score, trivial):
    if score >= 0.85:
        return "일치"
    if score >= 0.60:
        return "부분일치"
    return "상투문구(문면대조 생략)" if trivial else "상이·확인불가"


def match_all(recs, by_led):
    results = []
    for r in recs:
        key = (NFC(r["기관군"]), NFC(r["관계기관"]), NFC(r["소관부서"]))
        cands = by_led.get(key, [])
        row = dict(연번=r["연번"], 기관군=r["기관군"], 관계기관=r["관계기관"], 소관부서=r["소관부서"],
                   구분=r["구분"], 반영여부=r["반영여부"], 지문=norm(r["협의의견"])[:20])
        op_n = norm(r["협의의견"])
        if r["반영여부"] == "미회신":
            row.update(대조결과="미회신(원문 없음)", 유사도=None, 근거=[], 발췌="")
        elif not cands:
            row.update(대조결과="원문없음", 유사도=None, 근거=[], 발췌="")
        elif not op_n:
            row.update(대조결과="의견 공란", 유사도=None, 근거=[], 발췌="")
        else:
            scored = sorted(((coverage(op_n, d["_grams"]), d) for d in cands), key=lambda x: -x[0])
            best_s, best_d = scored[0]
            row.update(대조결과=verdict_of(best_s, bool(TRIV.match(r["협의의견"]))),
                       유사도=round(best_s, 3),
                       발췌=excerpt(op_n, best_d) if best_s >= 0.35 else "",
                       근거=[dict(파일=Path(best_d["file"]).name, 경로=best_d["file"],
                                 문서번호=best_d.get("문서번호", ""), 일자=best_d.get("시행일자", ""),
                                 방식=best_d["추출방식"])])
        results.append(row)
    return results


THIN = Side(style="thin", color="BFBFBF")
BD = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def write_xlsx(results, recs):
    op_by_n = {r["연번"]: re.sub(r"\s+", " ", r["협의의견"])[:120] for r in recs}
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "건별대조"
    cols = ["연번", "기관군", "관계기관", "소관부서", "구분", "반영여부", "대조결과", "유사도",
            "근거파일", "문서번호", "시행일자", "추출방식", "원문 발췌(최장 일치 대목)", "협의의견(참고)"]
    ws.append(cols)
    for i, w in enumerate([6, 9, 14, 16, 8, 9, 18, 8, 30, 16, 11, 10, 55, 55]):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i + 1)].width = w
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", size=10)
        c.fill = PatternFill("solid", fgColor="2E5C8A")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r in results:
        ev = r["근거"][0] if r["근거"] else {}
        ws.append([r["연번"], r["기관군"], r["관계기관"], r["소관부서"], r["구분"], r["반영여부"],
                   r["대조결과"], r["유사도"], ev.get("파일", ""), ev.get("문서번호", ""),
                   ev.get("일자", ""), ev.get("방식", ""), r["발췌"], op_by_n.get(r["연번"], "")])
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.border = BD
            c.alignment = Alignment(vertical="center", wrap_text=c.column_letter in ("I", "M", "N"))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws2 = wb.create_sheet("요약")
    ws2.append(["대조결과", "건수"])
    for k, v in Counter(r["대조결과"] for r in results).most_common():
        ws2.append([k, v])
    ws2.append(["합계", len(results)])
    wb.save(cfg.out("match_xlsx"))


def write_json(results):
    slim = {str(r["연번"]): dict(v=r["대조결과"], s=r["유사도"], fp=r["지문"],
                                ev=[dict(f=e["파일"], p=e["경로"], d=e["문서번호"], dt=e["일자"])
                                    for e in r["근거"]])
            for r in results}
    json.dump(slim, open(cfg.work("match_results.json"), "w", encoding="utf-8"), ensure_ascii=False)


if __name__ == "__main__":
    recs = json.load(open(cfg.work("records.json"), encoding="utf-8"))
    by_led = load_docs()
    results = match_all(recs, by_led)
    write_xlsx(results, recs)
    write_json(results)
    print(f"건별 대조 {len(results)}건 · 원문 {sum(len(v) for v in by_led.values())}건")
    for k, v in Counter(r["대조결과"] for r in results).most_common():
        print(f"  {k}: {v}")
    print("저장:", cfg.out("match_xlsx").relative_to(cfg.ROOT), "·", cfg.work("match_results.json").relative_to(cfg.ROOT))
